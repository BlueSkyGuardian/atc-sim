"""The real city pairs and airlines around the Agadir CTA (GMAC), from the workbook tools/traffic/
Agadir_CTA_Flight_Plan_Routes.xlsx (Flight Plan Database routes, Wikipedia airline lists), into
atcsim/data/flows/agadir_pairs.json for the scenario generator.

The workbook's routes are user-made plans, some against the RAD or one-way airways: they are not used to fly. Only
two things are taken from each one. First, the city pair and its airlines, with their group codes (easyJet EZY / EJU,
Ryanair RYR / MAY / RYS / RUK, Wizz WZZ / WMT, Transavia TRA / TVF). Second, whether the pair goes through GMAC or
passes nearby, measured here against the simulator's own sector polygons (the workbook's GMAC polygon was an
assumption, the southern part of the FIR). The generator flies each pair on our RAD-conformant flows.
  python tools/import_flight_pairs.py [workbook.xlsx]
"""
from __future__ import annotations

import json
import os
import re
import sys
import warnings

warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

import openpyxl  # noqa: E402

from atcsim import airspace as asp_mod  # noqa: E402
from atcsim.geometry import dist, point_in_polygon  # noqa: E402

SRC = os.path.join(HERE, "traffic", "Agadir_CTA_Flight_Plan_Routes.xlsx")
OUT = os.path.join(os.path.dirname(HERE), "atcsim", "data", "flows", "agadir_pairs.json")
NEAR_NM = 120.0          # passing within this of the GMAC sectors: nearby traffic (on the picture, not ours)

# pairs with no scheduled service, as the workbook found: replaced by the one that is flown (or dropped)
REPLACE = {("EGLL", "GCLP"): ("EGKK", "GCLP"), ("GCLP", "EGLL"): ("GCLP", "EGKK"),
           ("LEBL", "GCTS"): ("LEBL", "GCXO"), ("GCTS", "LEBL"): ("GCXO", "LEBL")}
REPLACE_OPS = {("EGKK", "GCLP"): ["BAW", "EZY", "TOM"], ("GCLP", "EGKK"): ["BAW", "EZY", "TOM"],
               ("LEBL", "GCXO"): ["VLG", "IBB"], ("GCXO", "LEBL"): ["VLG", "IBB"]}
NAMES = {"BAW": "British Airways", "EZY": "easyJet", "TOM": "TUI Airways", "VLG": "Vueling", "IBB": "Binter Canarias"}


def group_codes(icao: str, dep: str, dest: str) -> list:
    """[[code, share], ...]: the codes a group airline flies a route under."""
    ends = (dep, dest)
    uk = any(e.startswith(("EG", "EI")) for e in ends)
    if icao in ("EZY", "EJU", "EZS"):
        if any(e.startswith("LS") for e in ends):
            return [["EZS", 1.0]]
        return [["EZY", 1.0]] if uk else [["EJU", 1.0]]
    if icao == "RYR":
        return [["RYR", 0.7], ["RUK", 0.2], ["MAY", 0.1]] if uk else [["RYR", 0.75], ["MAY", 0.2], ["RYS", 0.05]]
    if icao == "WZZ":
        return [["WZZ", 0.7], ["WMT", 0.3]]
    if icao in ("TRA", "TVF"):
        return [["TVF", 1.0]] if any(e.startswith("LF") for e in ends) else [["TRA", 1.0]]
    return [[icao, 1.0]]


def main(src: str) -> None:
    wb = openpyxl.load_workbook(src)
    asp = asp_mod.get("agadir", ["ON", "WN", "WS", "OS"])
    P = asp.projection
    polys = {s.id: s.polygon for s in asp.sectors}
    casa = {s.id: s.polygon for s in asp_mod.get("agadir", ["N", "E", "C", "SS"]).sectors}
    ap = {a["id"]: (a["x"], a["y"]) for a in (asp.airports or []) if a.get("x") is not None}
    for d in asp.default_aerodromes:
        if "x" in d:
            ap.setdefault(d["id"], (d["x"], d["y"]))

    def coord(tok):
        m = re.match(r"^(\d{2})(\d{2})N(\d{1,4})-?(\d{1,3})(\d{2})W$", tok)
        return P.forward(int(m[1]) + int(m[2]) / 60, -(int(m[4]) + int(m[5]) / 60)) if m else None

    def edge_d(p, poly):
        best = 1e9
        for a, b in zip(poly, poly[1:] + poly[:1]):
            L2 = (b[0] - a[0]) ** 2 + (b[1] - a[1]) ** 2 or 1e-9
            t = max(0, min(1, ((p[0] - a[0]) * (b[0] - a[0]) + (p[1] - a[1]) * (b[1] - a[1])) / L2))
            best = min(best, dist(p, (a[0] + t * (b[0] - a[0]), a[1] + t * (b[1] - a[1]))))
        return best

    def where(dep, dest, names):
        pts = ([ap[dep]] if dep in ap else []) + [asp.waypoints[n] if n in asp.waypoints else coord(n) for n in names
                                                    if n in asp.waypoints or coord(n)] + ([ap[dest]] if dest in ap else [])
        secs, cs, dmin, nin = set(), set(), 1e9, 0.0
        for a, b in zip(pts, pts[1:]):
            n = max(1, int(dist(a, b) / 2))
            for k in range(n + 1):
                p = (a[0] + (b[0] - a[0]) * k / n, a[1] + (b[1] - a[1]) * k / n)
                hit = [s for s, poly in polys.items() if point_in_polygon(p, poly)]
                if hit:
                    secs.update(hit)
                    nin += dist(a, b) / (n + 1)
                else:
                    dmin = min(dmin, min(edge_d(p, poly) for poly in polys.values()))
                cs.update(s for s, poly in casa.items() if point_in_polygon(p, poly))
        return sorted(secs), round(nin), (0 if secs else round(dmin)), sorted(cs)

    routes = {}
    ws = wb["Routes"]
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        cat, dep, dest = r[0], r[1], r[2]
        if not dep or not dest:
            continue
        secs, nin, dmin, cs = where(dep, dest, (r[5] or "").split())
        routes[(dep, dest)] = {"category": cat, "gmac_sectors": secs, "gmac_nm": nin, "gmac_dist_nm": dmin,
                               "casablanca_sectors": cs, "workbook_status": r[7],
                               "cruise_ft": r[12] if isinstance(r[12], (int, float)) else None}
    for r in wb["Excluded (no GMAC)"].iter_rows(min_row=2, values_only=True):
        if r[1] and r[2] and (r[1], r[2]) not in routes:
            routes[(r[1], r[2])] = {"category": r[0], "gmac_sectors": [], "gmac_nm": 0, "gmac_dist_nm": None,
                                    "casablanca_sectors": [], "workbook_status": "excluded: " + str(r[3])}
    ops, no_svc = {}, set()
    ws = wb["Airlines by route"]
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        dep, dest, name, iata, icao, cs, svc, seen, note = r[:9]
        if not icao:
            if "no current scheduled service" in str(name or "").lower():
                no_svc.add((dep, dest))
            continue
        one = "one airport" in str(seen or "").lower()
        ops.setdefault((dep, dest), []).append({
            "icao": icao, "name": name, "callsign": cs, "service": svc, "codes": group_codes(icao, dep, dest),
            "weight": 0.5 if one else (0.6 if "charter" in str(svc or "").lower() or str(svc or "") == "Seasonal" else 1.0),
            **({"verify": "listed by one airport only"} if one else {}), **({"note": note} if note else {})})
    pairs, none = [], []
    for key in sorted(set(routes) | set(ops)):
        info = dict(routes.get(key, {}))
        dep, dest = key
        al = ops.get(key, [])
        if key in REPLACE:                         # no scheduled service: the pair that is flown instead
            none.append(list(key))
            dep, dest = REPLACE[key]
            info["replaces"] = f"{key[0]}-{key[1]} (no scheduled service)"
            al = [{"icao": c, "name": NAMES[c], "codes": group_codes(c, dep, dest), "weight": 1.0} for c in REPLACE_OPS[(dep, dest)]]
        if not al:
            if key in no_svc:
                none.append(list(key))             # no service at all (Paris - Bamako): never generated
            continue                               # (the excluded pairs: far from GMAC, no airlines read)
        pairs.append({"dep": dep, "dest": dest, **info, "airlines": al})
    doc = {"_note": "Real city pairs and their airlines around the Agadir CTA, imported by tools/import_flight_pairs.py "
                    "from tools/traffic/Agadir_CTA_Flight_Plan_Routes.xlsx. The generator draws a flight's departure, "
                    "destination and airline from these pairs when the flow it flies serves them (the flows keep "
                    "the RAD routes; the workbook's own routes are not flown). gmac_sectors / gmac_dist_nm: measured "
                    "on the workbook route against the simulator's sectors (nearby when within "
                    f"{int(NEAR_NM)} NM). codes: the airline's group codes and their shares. no_service: pairs not flown "
                    "(never generated).",
           "near_nm": NEAR_NM, "no_service": none, "pairs": pairs}
    open(OUT, "w").write(json.dumps(doc, indent=1, ensure_ascii=False) + "\n")
    through = sum(1 for p in pairs if p.get("gmac_sectors"))
    near = sum(1 for p in pairs if not p.get("gmac_sectors") and (p.get("gmac_dist_nm") or 1e9) <= NEAR_NM)
    print(f"{len(pairs)} pairs: {through} through GMAC, {near} nearby, {len(pairs) - through - near} further  -> {OUT}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else SRC)
